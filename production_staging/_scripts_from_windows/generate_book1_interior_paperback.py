import os
BUILD_ISBN = os.environ.get('BUILD_ISBN', 'ISBN-NOT-SET')
BUILD_IMPRINT = os.environ.get('BUILD_IMPRINT', 'Seventh City Press')
BUILD_AUTHOR = os.environ.get('BUILD_AUTHOR', 'Jason Carroll Holloway')
BUILD_TITLE = os.environ.get('BUILD_TITLE', '')
BUILD_EDITION_KEY = os.environ.get('BUILD_EDITION_KEY', '')
BUILD_OUTPUT = os.environ.get('BUILD_OUTPUT', None)
_w = os.environ.get('BUILD_WIDTH_PT')
_h = os.environ.get('BUILD_HEIGHT_PT')
BUILD_WIDTH_PT = float(_w) if _w else None
BUILD_HEIGHT_PT = float(_h) if _h else None
def format_isbn(raw):
    d = raw.replace('-','').replace(' ','')
    if len(d) == 13: return f"{d[0:3]}-{d[3]}-{d[4:9]}-{d[9:12]}-{d[12]}"
    return raw

"""
Masters X Book 1 Interior PDF Generator v5
Anathem-Inspired Typography · IngramSpark Upload

v5 fixes (2026-05-11, mirrored from Book 3 v5 build):
- Added a dedicated part-title page (P7) with "VOLUME ONE" eyebrow, no running
  head or folio, plus a quiet verso (P8) behind it.
- Suppressed running heads on every chapter opener; replaced the head with a
  drop folio at the bottom of the page (trade-publishing convention).
- Suppressed running head and folio entirely on display pages: half-title,
  blank versos, title page, copyright, epigraph, part-title, and the
  end-of-book closer.
- Added an "END OF VOLUME ONE - THE INHERITANCE OF FREQUENCY" closer page
  after the About-the-Author backmatter.
- Decimal-Hz frequency subtitles ("111.2 Hz - Prague", "107.3 Hz - Coda") now
  classify correctly so they render in chapter-subtitle style instead of as
  body text.
- Em-dash subtitle placeholders ("\u2014 \u00b7 <X>") fall through to the
  same Coda/subtitle style as a safety net for any unpatched DOCX.
- Two-pass build: first pass discovers chapter-opener page numbers, second
  pass renders the chrome with that set known.

v4 fixes (retained):
- Running header: "HOLLOWAY" on recto, "MASTERS X" on verso
- Italic run extraction from docx (preserves italic/bold formatting)
- Decorative section-break ornaments drawn as canvas graphics
"""

import sys, os, re
import docx
from reportlab.lib.units import inch
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.platypus import (
    Paragraph, Spacer, PageBreak, BaseDocTemplate, Frame, PageTemplate,
    NextPageTemplate
)
from reportlab.platypus.flowables import Flowable
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import Color
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from preview_source import chapter_one

# --- Paths ---
TRILOGY_DIR = Path(r"C:\Users\zh577\.openclaw\workspace\Corpus_Final_Export\The_Masters_Trilogy")
DOCX_PATH = Path(r"C:\Users\zh577\.gemini\antigravity\scratch\jasoncholloway\production_staging\_sources\build_docx\MASTERS_X_BOOK1_BUILD.docx")
OUTPUT_DIR = Path(r"C:\Users\zh577\.gemini\antigravity\scratch\jasoncholloway\Jason_Carroll_Holloway_Final_Export") #  / "Standard_BW_Paperback" / "Book1_The_Inheritance_of_Frequency"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_PDF = Path(BUILD_OUTPUT) if BUILD_OUTPUT else (OUTPUT_DIR / "INTERIOR_MASTERS_X_BOOK1_PB.pdf")

# Standard B&W paperback Bowker ISBN for Volume One.
BOOK_ISBN = (
    os.environ.get("MASTERSX_STANDARD_PB_BOOK1_ISBN", "").strip()
    or "979-8-2560-0804-8"
)

# --- Trim ---
TRIM_W = BUILD_WIDTH_PT if BUILD_WIDTH_PT else (5.5 * inch)
TRIM_H = BUILD_HEIGHT_PT if BUILD_HEIGHT_PT else (8.5 * inch)
M_GUTTER = 0.70 * inch
M_OUTSIDE = 0.674 * inch
M_TOP = 0.45 * inch
M_BOTTOM = 0.50 * inch
TEXT_W = TRIM_W - M_GUTTER - M_OUTSIDE
TEXT_H = TRIM_H - M_TOP - M_BOTTOM

# --- Fonts ---
pdfmetrics.registerFont(TTFont("Garamond", "C:/Windows/Fonts/GARA.TTF"))
pdfmetrics.registerFont(TTFont("GaramondBd", "C:/Windows/Fonts/GARABD.TTF"))
pdfmetrics.registerFont(TTFont("GaramondIt", "C:/Windows/Fonts/GARAIT.TTF"))
# Required for ReportLab <i>/<b> markup in Paragraphs. Without this family
# map, italic runs from DOCX are silently rendered as roman (body italics strip).
pdfmetrics.registerFontFamily(
    "Garamond",
    normal="Garamond",
    bold="GaramondBd",
    italic="GaramondIt",
    boldItalic="GaramondIt",
)

# --- Colors ---
MED = Color(0.4, 0.4, 0.4)
LT = Color(0.55, 0.55, 0.55)

# --- Styles ---
S_BODY = ParagraphStyle("Body", fontName="Garamond", fontSize=11, leading=15.5,
                         alignment=TA_JUSTIFY, firstLineIndent=18, spaceAfter=2.5, spaceBefore=0)
S_BODY1 = ParagraphStyle("Body1", parent=S_BODY, firstLineIndent=0)
S_CHNUM = ParagraphStyle("ChNum", fontName="Garamond", fontSize=11, leading=14,
                          alignment=TA_CENTER, textColor=MED, spaceAfter=8)
S_CHTITLE = ParagraphStyle("ChTitle", fontName="GaramondBd", fontSize=16, leading=20,
                            alignment=TA_CENTER, spaceAfter=6)
S_CHSUB = ParagraphStyle("ChSub", fontName="GaramondIt", fontSize=11, leading=14,
                          alignment=TA_CENTER, textColor=MED, spaceAfter=4)
S_HZ = ParagraphStyle("Hz", fontName="GaramondIt", fontSize=9, leading=12,
                       alignment=TA_CENTER, textColor=LT, spaceAfter=32)
S_HALF = ParagraphStyle("Half", fontName="Garamond", fontSize=20, leading=26, alignment=TA_CENTER)
S_TITLE = ParagraphStyle("Title", fontName="GaramondBd", fontSize=24, leading=30, alignment=TA_CENTER)
S_SUB = ParagraphStyle("Sub", fontName="GaramondIt", fontSize=14, leading=18,
                        alignment=TA_CENTER, textColor=MED)
S_AUTH = ParagraphStyle("Auth", fontName="Garamond", fontSize=14, leading=18, alignment=TA_CENTER)
S_COPY = ParagraphStyle("Copy", fontName="Garamond", fontSize=8, leading=11,
                         alignment=TA_CENTER, textColor=MED)
S_EPIG = ParagraphStyle("Epig", fontName="GaramondIt", fontSize=11, leading=15,
                         alignment=TA_CENTER, textColor=MED, leftIndent=36, rightIndent=36)
S_PREV = ParagraphStyle("Prev", fontName="GaramondIt", fontSize=12, leading=16,
                         alignment=TA_CENTER, textColor=MED, spaceAfter=8)
S_PARTEYE = ParagraphStyle("PartEye", fontName="Garamond", fontSize=11, leading=14,
                            alignment=TA_CENTER, textColor=LT, spaceAfter=18)
S_PARTTITLE = ParagraphStyle("PartTitle", fontName="GaramondBd", fontSize=22, leading=28,
                              alignment=TA_CENTER, spaceAfter=10)
S_PARTSUB = ParagraphStyle("PartSub", fontName="GaramondIt", fontSize=14, leading=18,
                            alignment=TA_CENTER, textColor=MED, spaceAfter=4)
S_CLOSER = ParagraphStyle("Closer", fontName="Garamond", fontSize=12, leading=16,
                           alignment=TA_CENTER, textColor=MED)


# --- Ornamental Rule (thin line) ---
class OrnRule(Flowable):
    def __init__(self, width):
        Flowable.__init__(self)
        self.width = width
        self.height = 8

    def draw(self):
        c = self.canv
        c.setStrokeColor(LT)
        c.setLineWidth(0.4)
        rw = self.width * 0.25
        x0 = (self.width - rw) / 2
        c.line(x0, 4, x0 + rw, 4)


# --- Decorative Section Break Ornament ---
class SectionOrnament(Flowable):
    """A decorative section break: thin rules flanking a central diamond cluster."""
    def __init__(self, width):
        Flowable.__init__(self)
        self.width = width
        self.height = 24

    def draw(self):
        c = self.canv
        mid_x = self.width / 2
        mid_y = self.height / 2

        c.setStrokeColor(MED)
        c.setFillColor(MED)

        diamond_size = 3.5
        c.saveState()
        p = c.beginPath()
        p.moveTo(mid_x, mid_y + diamond_size)
        p.lineTo(mid_x + diamond_size, mid_y)
        p.lineTo(mid_x, mid_y - diamond_size)
        p.lineTo(mid_x - diamond_size, mid_y)
        p.close()
        c.drawPath(p, fill=1, stroke=0)
        c.restoreState()

        sm = 2.2
        gap = 12
        for offset in [-gap, gap]:
            cx = mid_x + offset
            c.saveState()
            c.setLineWidth(0.5)
            p = c.beginPath()
            p.moveTo(cx, mid_y + sm)
            p.lineTo(cx + sm, mid_y)
            p.lineTo(cx, mid_y - sm)
            p.lineTo(cx - sm, mid_y)
            p.close()
            c.drawPath(p, fill=0, stroke=1)
            c.restoreState()

        rule_len = self.width * 0.1
        inner_edge = gap + sm + 8
        c.setLineWidth(0.3)
        c.line(mid_x - inner_edge - rule_len, mid_y, mid_x - inner_edge, mid_y)
        c.line(mid_x + inner_edge, mid_y, mid_x + inner_edge + rule_len, mid_y)


class ChapterMarker(Flowable):
    """Invisible marker for recto alignment + chapter-opener page tracking."""
    def __init__(self, chapter_name=""):
        Flowable.__init__(self)
        self.width = 0
        self.height = 0
        self.chapter_name = chapter_name

    def draw(self):
        pass

    def wrap(self, aW, aH):
        return (0, 0)


class NoChromeMarker(Flowable):
    """Invisible marker. Page on which this lands gets no header and no folio."""
    def __init__(self, label=""):
        Flowable.__init__(self)
        self.width = 0
        self.height = 0
        self.label = label

    def draw(self):
        pass

    def wrap(self, aW, aH):
        return (0, 0)


class InteriorDoc(BaseDocTemplate):
    """Two-pass-friendly document template."""

    FRONT_MATTER_PAGES = set(range(1, 7))  # P1..P6

    def __init__(
        self,
        filename,
        chapter_opener_pages=None,
        no_chrome_pages=None,
        **kw,
    ):
        self.chapter_name = ""
        self.total_pages = 0
        self.chapter_opener_pages = set(chapter_opener_pages or ())
        self.no_chrome_pages = set(no_chrome_pages or ()) | self.FRONT_MATTER_PAGES
        self.discovered_chapter_opener_pages = set()
        self.discovered_no_chrome_pages = set()
        BaseDocTemplate.__init__(self, filename, **kw)

        frame = Frame(
            M_GUTTER, M_BOTTOM, TEXT_W, TEXT_H,
            id="main",
            leftPadding=0, rightPadding=0,
            topPadding=0, bottomPadding=0,
        )
        self.addPageTemplates([
            PageTemplate("normal", frames=[frame], onPage=self._on_page),
        ])

    def _on_page(self, canvas, doc):
        pg = canvas.getPageNumber()
        self.total_pages = pg

        if pg in self.no_chrome_pages:
            return

        canvas.saveState()
        if pg in self.chapter_opener_pages:
            canvas.setFont("Garamond", 9)
            canvas.setFillColor(MED)
            canvas.drawCentredString(TRIM_W / 2, M_BOTTOM - 24, str(pg))
            canvas.restoreState()
            return

        hy = TRIM_H - M_TOP + 18
        canvas.setStrokeColor(LT)
        canvas.setLineWidth(0.3)
        canvas.line(M_GUTTER, hy - 4, TRIM_W - M_OUTSIDE, hy - 4)
        canvas.setFont("Garamond", 8)
        canvas.setFillColor(LT)
        if pg % 2 == 1:
            canvas.drawRightString(TRIM_W - M_OUTSIDE, hy, "HOLLOWAY")
        else:
            canvas.drawString(M_GUTTER, hy, "MASTERS X")
        canvas.setFont("Garamond", 9)
        canvas.setFillColor(MED)
        canvas.drawCentredString(TRIM_W / 2, M_BOTTOM - 24, str(pg))
        canvas.restoreState()

    def afterFlowable(self, flowable):
        if isinstance(flowable, ChapterMarker):
            self.chapter_name = flowable.chapter_name
            pg = self.canv.getPageNumber()
            if pg % 2 == 0:
                self.discovered_no_chrome_pages.add(pg)
                self.handle_pageBreak()
                pg = self.canv.getPageNumber()
            self.discovered_chapter_opener_pages.add(pg)
        elif isinstance(flowable, NoChromeMarker):
            pg = self.canv.getPageNumber()
            self.discovered_no_chrome_pages.add(pg)


# --- Extract from DOCX with formatting ---
# Ornament markers tolerated with either single OR double spaces between glyphs.
_CHBREAK_RE = re.compile(r"^\u25c7\s+\u25c6\s+\u25c7$")
_SECTBREAK_RE = re.compile(r"^\u2726\s+\u2295\s+\u2726$")


def classify(text):
    # Trilogy divider line that appears at the tail of each book in the
    # canonical DOCX (e.g. "END OF BOOK TWO: THE GRIMOIRE"). We render the
    # equivalent on a dedicated closer page, so suppress the body line.
    if text.upper().startswith("END OF BOOK"): return "skip"
    if _CHBREAK_RE.match(text): return "chbreak"
    if _SECTBREAK_RE.match(text): return "sectbreak"
    if text.startswith("CHAPTER "): return "chnum"
    if text in ("PROLOGUE", "EPILOGUE"): return "chnum"
    if text == "ONE YEAR LATER": return "episub"
    if text == "\u2014 \u00b7 Coda": return "coda"
    # Safety net: any em-dash subtitle placeholder routes through the coda flow.
    if re.match(r"^\u2014\s*\u00b7\s*.+$", text): return "coda"
    if text in ("SIX MONTHS LATER", "THREE WEEKS LATER"): return "timemark"
    # Frequency subtitle now also accepts decimal Hz values ("111.2 Hz - Prague").
    if re.match(r"^\d+(?:\.\d+)? Hz \u00b7 .+$", text): return "hz"
    if text == "MASTERS X": return "skip"
    if text == "MASTERS X: THE INHERITANCE OF FREQUENCY": return "skip"
    # The canonical DOCX's first paragraph is a readthrough header that mixes
    # the series title with an ornament glyph cluster, e.g.
    # "MASTERS X \u25c7 \u25c6 \u25c7 Trilogy Readthrough (Final)". Suppress
    # any line that names the trilogy readthrough so it doesn't leak onto the
    # blank verso behind the part-title.
    if "Trilogy Readthrough" in text: return "skip"
    if text.startswith("The Inheritance of Frequency"): return "skip"
    if text == "Jason Carroll Holloway": return "skip"
    if text.isupper() and len(text) < 60 and not text.startswith("CHAPTER"): return "chtitle"
    return "body"


def escape(text):
    text = text.replace("&", "&amp;")
    text = text.replace("<", "&lt;")
    text = text.replace(">", "&gt;")
    return text


def extract_with_formatting(docx_path, start=0, end=None):
    """Extract paragraphs with run-level italic/bold formatting."""
    doc = docx.Document(str(docx_path))
    out = []
    for i, p in enumerate(doc.paragraphs[start:end], start=start):
        plain = p.text.strip()
        if not plain:
            continue
        if "\ufffd" in plain:
            plain = plain.replace("\ufffd", "\u2014")

        cls = classify(plain)
        if cls == "coda" and out and out[-1]["c"] == "coda" and out[-1]["t"] == plain:
            continue

        if cls in ("chbreak", "sectbreak", "chnum", "chtitle", "hz",
                    "episub", "coda", "timemark", "skip"):
            out.append({"i": i, "t": plain, "c": cls, "rich": escape(plain)})
            continue

        rich_parts = []
        for run in p.runs:
            rt = run.text
            if not rt:
                continue
            rt = escape(rt)
            if run.italic and run.bold:
                rt = f"<b><i>{rt}</i></b>"
            elif run.italic:
                rt = f"<i>{rt}</i>"
            elif run.bold:
                rt = f"<b>{rt}</b>"
            rich_parts.append(rt)

        rich = "".join(rich_parts) if rich_parts else escape(plain)
        out.append({"i": i, "t": plain, "c": cls, "rich": rich})

    return out


# --- Build elements ---
def frontmatter():
    el = []
    # P1 recto: half-title
    el.append(Spacer(1, TEXT_H * 0.38))
    el.append(Paragraph("MASTERS X", S_HALF))
    el.append(PageBreak())
    # P2 verso: blank
    el.append(Spacer(1, 12))
    el.append(PageBreak())
    # P3 recto: full title
    el.append(Spacer(1, TEXT_H * 0.28))
    el.append(Paragraph("MASTERS X", S_TITLE))
    el.append(Spacer(1, 14))
    el.append(OrnRule(TEXT_W))
    el.append(Spacer(1, 10))
    el.append(Paragraph("The Inheritance of Frequency", S_SUB))
    el.append(Spacer(1, 4))
    el.append(Paragraph("Volume One", ParagraphStyle("V", parent=S_SUB, fontSize=11, textColor=LT)))
    el.append(Spacer(1, TEXT_H * 0.18))
    el.append(Paragraph("Jason Carroll Holloway", S_AUTH))
    el.append(PageBreak())
    # P4 verso: copyright
    el.append(Spacer(1, TEXT_H * 0.35))
    isbn_line = "ISBN " + format_isbn(BUILD_ISBN) if BUILD_ISBN and BUILD_ISBN != "ISBN-NOT-SET" else "ISBN " + BOOK_ISBN
    el.append(Paragraph(
        "Copyright \u00a9 2026 Jason Carroll Holloway. All rights reserved.<br/><br/>"
        "No part of this publication may be reproduced, distributed, or transmitted "
        "in any form or by any means without the prior written permission of the author.<br/><br/>"
        "This is a work of fiction. Names, characters, places, and incidents either are the "
        "product of the author\u2019s imagination or are used fictitiously.<br/><br/>"
        "Certain real locations, institutions, and public spaces are mentioned for atmospheric "
        "purposes. All events, characters, and interpretations associated with these locations "
        "are entirely fictional.<br/><br/>"
        "Published in the United States of America<br/><br/>"
        "First Edition<br/><br/>"
        + isbn_line,
        S_COPY))
    el.append(PageBreak())
    # P5 recto: epigraph
    el.append(Spacer(1, TEXT_H * 0.32))
    el.append(Paragraph(
        "\u201cThere\u2019s no chaos in nature, Blake.<br/>"
        "Only patterns we don\u2019t understand yet.\u201d", S_EPIG))
    el.append(Spacer(1, 14))
    el.append(Paragraph("\u2014 William Masters",
                         ParagraphStyle("EA", parent=S_EPIG, fontSize=9)))
    el.append(PageBreak())
    # P6 verso: blank
    el.append(NoChromeMarker("front_blank_p6"))
    el.append(Spacer(1, 12))
    el.append(PageBreak())
    # P7 recto: part-title (display page; no head, no folio)
    el.append(NoChromeMarker("part_title"))
    el.append(Spacer(1, TEXT_H * 0.30))
    el.append(Paragraph("VOLUME ONE", S_PARTEYE))
    el.append(OrnRule(TEXT_W))
    el.append(Spacer(1, 18))
    el.append(Paragraph("MASTERS X", S_PARTTITLE))
    el.append(Spacer(1, 4))
    el.append(Paragraph("The Inheritance of Frequency", S_PARTSUB))
    el.append(Spacer(1, 18))
    el.append(OrnRule(TEXT_W))
    el.append(PageBreak())
    # P8 verso: blank behind part-title
    el.append(NoChromeMarker("part_title_verso"))
    el.append(Spacer(1, 12))
    return el


def chapter_open(num_text, title, hz=None, sub=None):
    el = []
    el.append(Spacer(1, TEXT_H * 0.42))
    el.append(OrnRule(TEXT_W))
    el.append(Spacer(1, 14))
    el.append(Paragraph(num_text, S_CHNUM))
    if title:
        el.append(Paragraph(title, S_CHTITLE))
    if sub:
        el.append(Paragraph(sub, S_CHSUB))
    if hz:
        el.append(Spacer(1, 4))
        el.append(Paragraph(hz, S_HZ))
    el.append(Spacer(1, 10))
    el.append(OrnRule(TEXT_W))
    el.append(Spacer(1, 20))
    return el


def sect_break():
    return [
        Spacer(1, 20),
        SectionOrnament(TEXT_W),
        Spacer(1, 20),
    ]


def build_body(paras):
    el = []
    first = True
    i = 0
    while i < len(paras):
        p = paras[i]
        c = p["c"]

        if c in ("skip", "chbreak"):
            i += 1
            continue

        if c == "chnum":
            chnum = p["t"]
            chtitle = ""
            hz = None
            sub = None
            j = i + 1
            while j < len(paras):
                nc = paras[j]["c"]
                if nc == "chtitle":
                    chtitle = paras[j]["t"]
                    j += 1
                elif nc == "hz":
                    hz = paras[j]["t"]
                    j += 1
                elif nc == "episub":
                    sub = paras[j]["t"]
                    j += 1
                elif nc == "coda":
                    sub = (sub + "  " + paras[j]["t"]) if sub else paras[j]["t"]
                    j += 1
                else:
                    break

            # The Prologue is set in Bohemia 1267; the 109 Hz Kansas City tag
            # attached in the source DOCX is a paste artifact. Suppress.
            if chnum == "PROLOGUE":
                hz = None

            el.append(PageBreak())
            el.append(ChapterMarker(chtitle or chnum))
            el.extend(chapter_open(chnum, chtitle, hz, sub))
            first = True
            i = j
            continue

        if c == "sectbreak":
            # Suppress a sectbreak that's the last thing in its chapter.
            k = i + 1
            while k < len(paras) and paras[k]["c"] in ("chbreak", "sectbreak", "skip"):
                k += 1
            if k >= len(paras) or paras[k]["c"] == "chnum":
                i += 1
                continue
            el.extend(sect_break())
            first = True
            i += 1
            continue

        if c == "timemark":
            el.append(PageBreak())
            el.append(ChapterMarker(p["t"]))
            el.extend(chapter_open(p["t"], ""))
            first = True
            i += 1
            continue

        rich = p["rich"]
        el.append(Paragraph(rich, S_BODY1 if first else S_BODY))
        first = False
        i += 1
    return el


def preview_section(paras):
    # Without body text the scaffolding below renders a chapter opener followed
    # immediately by "End of Preview" — an empty teaser. Emit nothing instead.
    if not paras:
        return []
    el = []
    el.append(PageBreak())
    el.append(ChapterMarker("PREVIEW"))
    el.append(Spacer(1, TEXT_H * 0.32))
    el.append(Paragraph("A PREVIEW OF", ParagraphStyle("PL", parent=S_PREV, fontSize=11)))
    el.append(Spacer(1, 8))
    el.append(Paragraph("MASTERS X", S_TITLE))
    el.append(Spacer(1, 6))
    el.append(Paragraph("The Grimoire", S_SUB))
    el.append(Spacer(1, 4))
    el.append(Paragraph("Volume Two", ParagraphStyle("V2", parent=S_SUB, fontSize=11, textColor=LT)))
    el.append(Spacer(1, 28))
    el.append(OrnRule(TEXT_W))
    el.append(Spacer(1, 14))
    el.append(Paragraph("Available now", ParagraphStyle("AN", parent=S_PREV, fontSize=10)))
    el.append(PageBreak())
    el.append(Spacer(1, 12))
    el.append(PageBreak())
    el.append(ChapterMarker("THE STONE COTTAGE"))
    el.extend(chapter_open("CHAPTER ONE", "THE STONE COTTAGE", "3.915 Hz \u00b7 Iceland"))
    first = True
    for p in paras:
        c = p["c"]
        if c in ("chnum", "chtitle", "hz", "chbreak", "skip"):
            continue
        if c == "sectbreak":
            el.extend(sect_break())
            first = True
            continue
        rich = p["rich"]
        el.append(Paragraph(rich, S_BODY1 if first else S_BODY))
        first = False
    el.append(Spacer(1, 40))
    el.append(OrnRule(TEXT_W))
    el.append(Spacer(1, 14))
    el.append(Paragraph("End of Preview", ParagraphStyle("EP", parent=S_PREV, fontSize=10)))
    el.append(Spacer(1, 8))
    el.append(Paragraph("The Grimoire continues at your bookseller.",
                          ParagraphStyle("C2", parent=S_PREV, fontSize=10, fontName="Garamond")))
    return el


def backmatter():
    el = []
    el.append(PageBreak())
    el.append(ChapterMarker("ABOUT THE AUTHOR"))
    el.append(Spacer(1, TEXT_H * 0.20))
    el.append(Paragraph("ABOUT THE AUTHOR", ParagraphStyle("AH", parent=S_CHTITLE, fontSize=14)))
    el.append(Spacer(1, 14))
    el.append(OrnRule(TEXT_W))
    el.append(Spacer(1, 20))
    el.append(Paragraph(
        "Jason Carroll Holloway holds an M.A. in English Literature from Mercy University in Dobbs Ferry, New York, as well as various degrees and certificates in Psychology, Sociology, Creative Writing, and Data Analytics. He is a writer and researcher whose work explores the intersection of acoustic science, medieval scholarship, and human consciousness. He lives in Kansas City.",
        ParagraphStyle("AB", parent=S_BODY1, alignment=TA_CENTER, fontSize=11, leading=15)))
    el.append(Spacer(1, 36))
    el.append(Paragraph("ALSO BY JASON CARROLL HOLLOWAY",
                          ParagraphStyle("ABH", parent=S_CHTITLE, fontSize=12)))
    el.append(Spacer(1, 14))
    el.append(OrnRule(TEXT_W))
    el.append(Spacer(1, 20))
    for t in [
        "Masters X: The Grimoire (Volume Two)",
        "Masters X: The Kingdom (Volume Three)",
        "Innocence, Desire, and the Architecture of the Fall",
    ]:
        el.append(Paragraph(f"<i>{escape(t)}</i>",
                              ParagraphStyle("AI", parent=S_BODY1, alignment=TA_CENTER,
                                              fontSize=11, leading=18, spaceAfter=4)))
    return el


def closer_page():
    """End-of-book closer. Recto display page with no head or folio."""
    el = []
    el.append(PageBreak())
    el.append(ChapterMarker("CLOSER"))
    el.append(NoChromeMarker("end_of_book"))
    el.append(Spacer(1, TEXT_H * 0.42))
    el.append(Paragraph("END OF VOLUME ONE", S_PARTEYE))
    el.append(OrnRule(TEXT_W))
    el.append(Spacer(1, 14))
    el.append(Paragraph("MASTERS X", S_PARTTITLE))
    el.append(Spacer(1, 4))
    el.append(Paragraph("The Inheritance of Frequency", S_PARTSUB))
    el.append(Spacer(1, 14))
    el.append(OrnRule(TEXT_W))
    return el


# --- Main ---
def _build_elements(book1, preview):
    elements = []
    elements.extend(frontmatter())
    elements.extend(build_body(book1))
    elements.extend(preview_section(preview))
    elements.extend(backmatter())
    elements.extend(closer_page())
    return elements


def main():
    print("=" * 60)
    print("MASTERS X \u00b7 BOOK 1 \u00b7 INTERIOR PDF GENERATOR v5")
    print("Anathem Register \u00b7 IngramSpark Upload")
    print("=" * 60)

    print("\n[1/5] Extracting Book 1 with formatting...")
    book1 = extract_with_formatting(DOCX_PATH, 0, None)
    italic_count = sum(1 for p in book1 if "<i>" in p["rich"])
    print(f"      {len(book1)} paragraphs ({italic_count} with italic runs)")

    print("[2/5] Loading Book 2 preview from canonical build DOCX...")
    preview = chapter_one(1)
    print(f"      {len(preview)} paragraphs, {sum(len(p['t'].split()) for p in preview)} words")

    print("[3/5] Pass 1: discovering layout (chapter openers, display pages)...")
    scratch = OUTPUT_DIR / "_pass1_scratch.pdf"
    doc1 = InteriorDoc(
        str(scratch),
        pagesize=(TRIM_W, TRIM_H),
        title="Masters X: The Inheritance of Frequency",
        author="Jason Carroll Holloway",
    )
    doc1.build(_build_elements(book1, preview))
    chapter_openers = set(doc1.discovered_chapter_opener_pages)
    no_chrome = set(doc1.discovered_no_chrome_pages)
    print(f"      chapter openers: {sorted(chapter_openers)}")
    print(f"      no-chrome pages: {sorted(no_chrome)}")
    try:
        scratch.unlink()
    except OSError:
        pass

    print("[4/5] Pass 2: rendering final PDF...")
    doc = InteriorDoc(
        str(OUTPUT_PDF),
        pagesize=(TRIM_W, TRIM_H),
        title="Masters X: The Inheritance of Frequency",
        author="Jason Carroll Holloway",
        chapter_opener_pages=chapter_openers,
        no_chrome_pages=no_chrome,
    )
    doc.build(_build_elements(book1, preview))

    print(f"\n[5/5] Complete!")
    print(f"      Output: {OUTPUT_PDF}")
    print(f"      Pages: {doc.total_pages}")
    print(f"\n{'=' * 60}")
    print(f"TOTAL PAGES: {doc.total_pages}")
    print(f"{'=' * 60}")


if __name__ == "__main__":
    main()
